"""app/excel_engine.py — Motor de actualización Excel (COM / openpyxl fallback)"""
from __future__ import annotations

import os
import sys
import time
import logging
import threading
import traceback
from pathlib import Path
from dataclasses import dataclass, field


# ── win32com opcional ─────────────────────────────────────────────────────────
try:
    import win32com.client
    import pythoncom
    import pywintypes
    WIN32_AVAILABLE = True
except ImportError:
    WIN32_AVAILABLE = False

# HRESULT devuelto cuando Excel está ocupado y rechaza la llamada COM
_COM_CALL_REJECTED = -2147418111   # RPC_E_CALL_REJECTED
_COM_RETRY_ATTEMPTS = 6
_COM_RETRY_WAIT = 3  # segundos entre reintentos

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


# ══════════════════════════════════════════════════════════════════════════════
@dataclass
class EngineResult:
    success: bool = False
    connections_found: int = 0
    pivots_ok: int = 0
    pivots_err: int = 0
    error_msg: str = ""
    duration_s: float = 0.0
    pivots_completed: list = field(default_factory=list)  # [{"sheet": .., "pivot": ..}, ...]


@dataclass
class EngineConfig:
    file_path: str
    refresh_connections: bool = True
    refresh_pivots: bool = True
    save_on_success: bool = True
    excel_visible: bool = False
    lock_timeout: int = 120
    lock_retry: int = 5
    lock_max_retries: int = 5
    refresh_timeout: int = 300
    refresh_check: int = 3
    stop_event: threading.Event = field(default_factory=threading.Event)
    pivot_max_retries: int = 3          # reintentos intra-ejecución por pivot
    pivot_retry_delay_s: int = 60       # espera (s) entre reintentos de pivot
    skip_pivots: set = field(default_factory=set)  # set de (sheet_name, pivot_name) a saltar


# ══════════════════════════════════════════════════════════════════════════════
# UTILIDADES
# ══════════════════════════════════════════════════════════════════════════════

# Atributos de OneDrive Files On-Demand (placeholders en la nube)
_ONEDRIVE_CLOUD_ATTRS = 0x400000 | 0x40000  # RECALL_ON_DATA_ACCESS | RECALL_ON_OPEN


def resolve_path(raw: str) -> str:
    """Expande variables de entorno y normaliza la ruta sin cambiar el disco."""
    expanded = os.path.expandvars(raw.strip())
    p = Path(expanded)
    # Si la ruta es absoluta, usar os.path.normpath para no depender del CWD
    # (Path.resolve() en D: podría mutar rutas de C: en algunos entornos)
    if p.is_absolute():
        return os.path.normpath(str(p))
    # Si es relativa, resolver desde el directorio de trabajo del proceso
    return os.path.normpath(os.path.abspath(expanded))


def _is_onedrive_placeholder(path: str) -> bool:
    """Devuelve True si el archivo es un placeholder OneDrive (no descargado)."""
    if sys.platform != "win32":
        return False
    try:
        import ctypes
        attrs = ctypes.windll.kernel32.GetFileAttributesW(path)
        return attrs != -1 and bool(attrs & _ONEDRIVE_CLOUD_ATTRS)
    except Exception:
        return False


def _trigger_onedrive_download(path: str, log: logging.Logger, timeout: int = 120):
    """
    Fuerza la descarga del placeholder OneDrive y bloquea hasta que el archivo
    esté físicamente disponible (o se agote el timeout).
    """
    import ctypes
    import subprocess

    # Disparar hidratación abriendo en lectura (Cloud Provider IRP)
    try:
        with open(path, "rb") as f:
            f.read(1)
    except Exception:
        pass

    # Marcar como "pinned" para que OneDrive lo mantenga local
    try:
        subprocess.run(["attrib", "+P", "-U", path],
                       capture_output=True, timeout=10)
    except Exception:
        pass

    log.info("Descarga de OneDrive iniciada. Esperando sincronización...")

    RECALL_FLAGS = 0x400000 | 0x40000  # RECALL_ON_DATA_ACCESS | RECALL_ON_OPEN
    deadline = time.time() + timeout
    while time.time() < deadline:
        attrs = ctypes.windll.kernel32.GetFileAttributesW(path)
        if attrs != -1 and not (attrs & RECALL_FLAGS):
            log.info("Archivo OneDrive hidratado correctamente.")
            return
        time.sleep(2)

    log.warning(f"Timeout ({timeout}s) esperando descarga de OneDrive. Se intentará continuar.")


def _read_laccdb_users(accdb_path: str) -> list[tuple[str, str]]:
    """Lee el .laccdb hermano de un .accdb y devuelve [(computer, user), ...].

    El formato es binario pero estable: bloques de 64 bytes con 32 bytes para
    computer name y 32 bytes para username, rellenados con NUL. El archivo es
    legible aún cuando Access lo tiene abierto.
    """
    p = Path(accdb_path)
    if p.suffix.lower() not in (".accdb", ".mdb"):
        return []
    lock = p.with_suffix(".laccdb" if p.suffix.lower() == ".accdb" else ".ldb")
    if not lock.exists():
        return []
    try:
        data = lock.read_bytes()
    except OSError:
        return []
    entries: list[tuple[str, str]] = []
    for i in range(0, len(data), 64):
        block = data[i:i + 64]
        if len(block) < 64:
            break
        computer = block[0:32].split(b"\x00", 1)[0].decode("latin-1", "ignore").strip()
        user = block[32:64].split(b"\x00", 1)[0].decode("latin-1", "ignore").strip()
        if computer or user:
            entries.append((computer, user))
    # Deduplicar manteniendo orden
    seen: set[tuple[str, str]] = set()
    unique: list[tuple[str, str]] = []
    for e in entries:
        if e not in seen:
            seen.add(e)
            unique.append(e)
    return unique


def _read_xlsx_lock_owner(xlsx_path: str) -> str | None:
    """Lee el ~$archivo.xlsx y devuelve el usuario que tiene el libro abierto.

    Formato: byte 0 = longitud del nombre (en chars), bytes 1.. = nombre en
    UTF-16-LE (Excel trunca a 15 chars). Devuelve None si no se puede leer.
    """
    p = Path(xlsx_path)
    lock = p.parent / f"~${p.name}"
    if not lock.exists():
        return None
    try:
        data = lock.read_bytes()
    except OSError:
        return None
    if len(data) < 3:
        return None
    try:
        length = data[0]
        raw = data[1:1 + length * 2]
        name = raw.decode("utf-16-le", "ignore").rstrip("\x00").strip()
        return name or None
    except Exception:
        return None


def _rm_get_locking_processes(path: str) -> list[dict]:
    """Vía Windows Restart Manager API: lista procesos que tienen handles abiertos
    sobre `path`. Devuelve [{pid, name, service, session_id, start_time}, ...]
    o [] si la API no está disponible o no hay bloqueos visibles. No requiere admin.
    """
    if sys.platform != "win32":
        return []
    try:
        import ctypes
        from ctypes import wintypes
    except Exception:
        return []
    try:
        rstrtmgr = ctypes.WinDLL("rstrtmgr")
    except OSError:
        return []

    CCH_RM_SESSION_KEY = 32
    CCH_RM_MAX_APP_NAME = 255
    CCH_RM_MAX_SVC_NAME = 63
    ERROR_SUCCESS = 0
    ERROR_MORE_DATA = 234

    class RM_UNIQUE_PROCESS(ctypes.Structure):
        _fields_ = [
            ("dwProcessId", wintypes.DWORD),
            ("ProcessStartTime", wintypes.FILETIME),
        ]

    class RM_PROCESS_INFO(ctypes.Structure):
        _fields_ = [
            ("Process", RM_UNIQUE_PROCESS),
            ("strAppName", ctypes.c_wchar * (CCH_RM_MAX_APP_NAME + 1)),
            ("strServiceShortName", ctypes.c_wchar * (CCH_RM_MAX_SVC_NAME + 1)),
            ("ApplicationType", ctypes.c_int),
            ("AppStatus", ctypes.c_ulong),
            ("TSSessionId", wintypes.DWORD),
            ("bRestartable", wintypes.BOOL),
        ]

    rstrtmgr.RmStartSession.argtypes = [
        ctypes.POINTER(wintypes.DWORD), wintypes.DWORD, ctypes.c_wchar_p,
    ]
    rstrtmgr.RmStartSession.restype = wintypes.DWORD
    rstrtmgr.RmRegisterResources.argtypes = [
        wintypes.DWORD, wintypes.UINT, ctypes.POINTER(ctypes.c_wchar_p),
        wintypes.UINT, ctypes.c_void_p, wintypes.UINT, ctypes.POINTER(ctypes.c_wchar_p),
    ]
    rstrtmgr.RmRegisterResources.restype = wintypes.DWORD
    rstrtmgr.RmGetList.argtypes = [
        wintypes.DWORD, ctypes.POINTER(wintypes.UINT), ctypes.POINTER(wintypes.UINT),
        ctypes.POINTER(RM_PROCESS_INFO), ctypes.POINTER(wintypes.DWORD),
    ]
    rstrtmgr.RmGetList.restype = wintypes.DWORD
    rstrtmgr.RmEndSession.argtypes = [wintypes.DWORD]
    rstrtmgr.RmEndSession.restype = wintypes.DWORD

    session = wintypes.DWORD()
    session_key = (ctypes.c_wchar * (CCH_RM_SESSION_KEY + 1))()
    rc = rstrtmgr.RmStartSession(ctypes.byref(session), 0, session_key)
    if rc != ERROR_SUCCESS:
        return []

    try:
        files_arr = (ctypes.c_wchar_p * 1)(path)
        rc = rstrtmgr.RmRegisterResources(
            session, 1, files_arr, 0, None, 0, None,
        )
        if rc != ERROR_SUCCESS:
            return []

        proc_needed = wintypes.UINT(0)
        proc_count = wintypes.UINT(0)
        reasons = wintypes.DWORD(0)

        rc = rstrtmgr.RmGetList(
            session, ctypes.byref(proc_needed), ctypes.byref(proc_count),
            None, ctypes.byref(reasons),
        )
        if rc == ERROR_SUCCESS and proc_needed.value == 0:
            return []
        if rc != ERROR_MORE_DATA:
            return []

        n = proc_needed.value
        procs = (RM_PROCESS_INFO * n)()
        proc_count = wintypes.UINT(n)
        rc = rstrtmgr.RmGetList(
            session, ctypes.byref(proc_needed), ctypes.byref(proc_count),
            procs, ctypes.byref(reasons),
        )
        if rc != ERROR_SUCCESS:
            return []

        from datetime import datetime
        results: list[dict] = []
        epoch_diff_100ns = 116444736000000000  # 1601-01-01 → 1970-01-01
        for i in range(proc_count.value):
            entry = procs[i]
            ft = entry.Process.ProcessStartTime
            start_dt = None
            try:
                ts100ns = (ft.dwHighDateTime << 32) | ft.dwLowDateTime
                if ts100ns > 0:
                    start_dt = datetime.fromtimestamp(
                        (ts100ns - epoch_diff_100ns) / 10_000_000
                    )
            except Exception:
                pass
            results.append({
                "pid": entry.Process.dwProcessId,
                "name": entry.strAppName,
                "service": entry.strServiceShortName,
                "session_id": entry.TSSessionId,
                "start_time": start_dt,
            })
        return results
    finally:
        try:
            rstrtmgr.RmEndSession(session)
        except Exception:
            pass


def _find_orphan_runs_for_file(path: str, min_age_seconds: int = 120) -> list[dict]:
    """Busca RunLogs en estado 'running' que toquen `path` y lleven al menos
    `min_age_seconds` corriendo. Sync sqlite — pensado para llamarse desde
    el thread sync de wait_for_file. Devuelve [] si no hay DB o falla."""
    try:
        from datetime import datetime, timedelta
        from app.config import settings
        import sqlite3
        db_file = Path(settings.db_path)
        if not db_file.exists():
            return []
        cutoff = (datetime.now() - timedelta(seconds=min_age_seconds)).isoformat()
        fname = Path(path).name
        # 'path' en pipeline_config aparece JSON-escapado, así que matcheamos por basename
        like_basename = f"%{fname}%"
        conn = sqlite3.connect(str(db_file), timeout=2)
        try:
            conn.row_factory = sqlite3.Row
            cur = conn.execute(
                """
                SELECT r.id AS run_id, r.task_id, r.task_name, r.started_at
                FROM run_logs r
                JOIN tasks t ON t.id = r.task_id
                WHERE r.status = 'running'
                  AND r.started_at < ?
                  AND (t.file_path = ? OR t.pipeline_config LIKE ?)
                ORDER BY r.started_at ASC
                LIMIT 5
                """,
                (cutoff, path, like_basename),
            )
            return [dict(row) for row in cur.fetchall()]
        finally:
            conn.close()
    except Exception:
        return []


def _lock_reason(path: str) -> str:
    """Detecta el motivo y autor probable del bloqueo. Combina:
    1) `.laccdb` parsing (Access)        → usuario + máquina
    2) `~$archivo.xlsx` parsing (Excel)  → usuario
    3) Windows Restart Manager API       → proceso, PID y hora de inicio
    4) RunLog 'running' antiguo en DB    → indica orfanato de un run previo
    """
    p = Path(path)
    parts: list[str] = []

    if _is_onedrive_placeholder(path):
        parts.append("archivo OneDrive en la nube (descargando)")

    # — Access (.laccdb) ————————————————————————————————————————————
    if p.suffix.lower() in (".accdb", ".mdb"):
        users = _read_laccdb_users(path)
        if users:
            shown = ", ".join(f"{u or '?'}@{c or '?'}" for c, u in users[:3])
            extra = f" (+{len(users)-3} más)" if len(users) > 3 else ""
            parts.append(f"Access lo tiene abierto: {shown}{extra}")

    # — Excel (~$) ————————————————————————————————————————————————
    excel_lock = p.parent / f"~${p.name}"
    if excel_lock.exists():
        owner = _read_xlsx_lock_owner(path)
        if owner:
            parts.append(f"Excel lo tiene abierto (~$lock, usuario: {owner})")
        else:
            parts.append("Excel lo tiene abierto (~$lock detectado)")

    # — Restart Manager (cualquier OS handle) ————————————————————
    rm_procs = _rm_get_locking_processes(path)
    for proc in rm_procs:
        start = proc.get("start_time")
        start_str = start.strftime("%Y-%m-%d %H:%M:%S") if start else "?"
        parts.append(
            f"proceso {proc.get('name') or '?'} "
            f"PID={proc.get('pid')} sesión={proc.get('session_id')} "
            f"iniciado={start_str}"
        )

    # — RunLogs 'running' antiguos ————————————————————————————————
    orphans = _find_orphan_runs_for_file(path)
    for orun in orphans:
        parts.append(
            f"run anterior huérfano: task='{orun.get('task_name')}' "
            f"run_id={orun.get('run_id')} iniciado={orun.get('started_at')}"
        )

    if not parts:
        tmp_files = list(p.parent.glob("*.tmp"))
        if tmp_files:
            return "posible sincronización de OneDrive en curso"
        return "proceso externo desconocido"

    return "; ".join(parts)


def wait_for_file(path: str, timeout: int, interval: int,
                  log: logging.Logger, max_retries: int = 0,
                  stop_event: threading.Event | None = None) -> bool:
    """
    Espera hasta que el archivo esté disponible.
    - timeout: tiempo máximo total en segundos.
    - max_retries: número máximo de intentos (0 = ilimitado, solo por timeout).
    - stop_event: si se activa, aborta la espera inmediatamente.
    """
    deadline = time.time() + timeout
    attempt = 0
    while time.time() < deadline:
        if stop_event and stop_event.is_set():
            log.warning("Ejecución detenida manualmente durante espera de archivo.")
            return False
        try:
            with open(path, "r+b"):
                return True
        except (PermissionError, OSError):
            attempt += 1
            elapsed = round(timeout - (deadline - time.time()))
            reason = _lock_reason(path)
            limit_info = f"/{max_retries}" if max_retries else ""
            log.warning(
                f"Archivo bloqueado (intento {attempt}{limit_info}, {elapsed}s/{timeout}s) "
                f"— {reason}. Reintentando en {interval}s…"
            )
            if max_retries and attempt >= max_retries:
                log.error(f"Se alcanzó el límite de {max_retries} reintentos. Tarea cancelada.")
                return False
            # Esperar en pequeños intervalos para responder rápido al stop_event
            end = time.time() + interval
            while time.time() < end:
                if stop_event and stop_event.is_set():
                    log.warning("Ejecución detenida manualmente durante espera de archivo.")
                    return False
                time.sleep(0.5)
    return False


# ══════════════════════════════════════════════════════════════════════════════
# MOTOR COM
# ══════════════════════════════════════════════════════════════════════════════

class ExcelCOMUpdater:
    def __init__(self, cfg: EngineConfig, log: logging.Logger):
        self.cfg = cfg
        self.log = log
        self.xl = None
        self.wb = None

    def open(self):
        pythoncom.CoInitialize()
        self.xl = win32com.client.DispatchEx("Excel.Application")
        self.xl.Visible = self.cfg.excel_visible
        self.xl.DisplayAlerts = False
        self.xl.AskToUpdateLinks = False
        self.wb = self.xl.Workbooks.Open(
            self.cfg.file_path, UpdateLinks=0, ReadOnly=False
        )
        # Esperar a que Excel termine de cargar el libro antes de seguir.
        # Sin este delay, el primer acceso a wb.Connections puede devolver
        # RPC_E_CALL_REJECTED (-2147418111) porque Excel aún no está listo.
        time.sleep(2)
        self.log.info("Libro abierto (COM).")

    def close(self, save: bool):
        if self.wb:
            try:
                if save:
                    self.wb.Save()
                    self.log.info("Guardado.")
                self.wb.Close(SaveChanges=False)
            except Exception as e:
                self.log.error(f"Error cerrando libro: {e}")
        if self.xl:
            try:
                self.xl.Quit()
            except Exception:
                pass
        pythoncom.CoUninitialize()

    def _is_connection_refreshing(self, conn) -> bool:
        """Comprueba si una conexión COM está actualizando. Maneja tipos no soportados."""
        for attr in ("ODBCConnection", "OLEDBConnection",
                     "TextConnection", "WorksheetDataConnection"):
            try:
                # getattr en COM puede lanzar excepción si el atributo no existe
                sub = getattr(conn, attr)
                if sub is not None and sub.Refreshing:
                    return True
            except Exception:
                pass
        return False

    def _com_retry(self, func, retries: int = 5, delay: float = 2.0):
        """Reintenta una llamada COM si Excel devuelve RPC_E_CALL_REJECTED."""
        import pywintypes
        RPC_E_CALL_REJECTED = -2147418111
        for attempt in range(retries):
            try:
                return func()
            except pywintypes.com_error as e:
                if e.args[0] == RPC_E_CALL_REJECTED and attempt < retries - 1:
                    self.log.debug(f"Excel ocupado (intento {attempt + 1}/{retries}), reintentando en {delay}s...")
                    time.sleep(delay)
                else:
                    raise

    def _refresh_connections(self) -> tuple[bool, int]:
        conns = self._com_retry(lambda: self.wb.Connections)
        n = conns.Count
        self.log.info(f"Conexiones: {n}")
        if n == 0:
            return True, 0

        self.wb.RefreshAll()
        deadline = time.time() + self.cfg.refresh_timeout
        while time.time() < deadline:
            if self.cfg.stop_event.is_set():
                self.log.warning("Ejecución detenida manualmente durante actualización de conexiones.")
                return False, n
            time.sleep(self.cfg.refresh_check)
            still_refreshing = []
            for i in range(1, n + 1):
                try:
                    conn = conns.Item(i)
                    if self._is_connection_refreshing(conn):
                        still_refreshing.append(conn.Name)
                except Exception as e:
                    self.log.debug(f"No se pudo leer estado de conexión {i}: {e}")
            if not still_refreshing:
                self.log.info("Todas las conexiones finalizadas.")
                return True, n
            self.log.debug(f"Esperando conexiones: {still_refreshing}")
        self.log.error("Timeout esperando conexiones.")
        return False, n

    def _com_retry(self, fn, *args):
        """Ejecuta fn(*args) reintentando si Excel rechaza la llamada (RPC_E_CALL_REJECTED)."""
        for attempt in range(_COM_RETRY_ATTEMPTS):
            try:
                return fn(*args)
            except pywintypes.com_error as e:
                if e.hresult == _COM_CALL_REJECTED and attempt < _COM_RETRY_ATTEMPTS - 1:
                    self.log.debug(
                        f"COM rechazado por Excel (intento {attempt + 1}/{_COM_RETRY_ATTEMPTS}), "
                        f"reintentando en {_COM_RETRY_WAIT}s..."
                    )
                    time.sleep(_COM_RETRY_WAIT)
                else:
                    raise

    def _refresh_pivots(self, res: "EngineResult") -> tuple[int, int]:
        ok = err = 0
        max_retries   = self.cfg.pivot_max_retries
        retry_delay   = self.cfg.pivot_retry_delay_s
        skip          = self.cfg.skip_pivots  # set de (sheet_name, pivot_name)

        for s in range(1, self.wb.Sheets.Count + 1):
            sheet      = self._com_retry(lambda idx=s: self.wb.Sheets(idx))
            sheet_name = self._com_retry(lambda sh=sheet: sh.Name)
            for p in range(1, self._com_retry(lambda sh=sheet: sh.PivotTables().Count) + 1):
                pt      = self._com_retry(lambda sh=sheet, idx=p: sh.PivotTables(idx))
                pt_name = self._com_retry(lambda t=pt: t.Name)

                # ── Skip: pivot ya completado en run anterior ─────────────
                if (sheet_name, pt_name) in skip:
                    self.log.info(f"  ↷ PivotTable '{pt_name}' en '{sheet_name}' — saltado (ya actualizado)")
                    continue

                # ── Retry intra-ejecución ─────────────────────────────────
                last_exc = None
                for attempt in range(max_retries + 1):
                    try:
                        self._com_retry(lambda t=pt: t.RefreshTable())
                        ok += 1
                        self.log.info(f"  ✔ PivotTable '{pt_name}' en '{sheet_name}'")
                        res.pivots_completed.append({"sheet": sheet_name, "pivot": pt_name})
                        last_exc = None
                        break
                    except Exception as e:
                        last_exc = e
                        if attempt < max_retries:
                            self.log.warning(
                                f"  ↻ Reintento {attempt + 1}/{max_retries} "
                                f"PivotTable '{pt_name}' en '{sheet_name}' "
                                f"(esperando {retry_delay}s)..."
                            )
                            time.sleep(retry_delay)
                        else:
                            err += 1
                            self.log.error(
                                f"  ✘ PivotTable '{pt_name}' en '{sheet_name}' "
                                f"(falló {max_retries + 1} intentos): {last_exc}"
                            )
                            # Pausa para que Excel vuelva a estado estable
                            time.sleep(2)
        return ok, err

    def run(self) -> EngineResult:
        t0 = time.time()
        res = EngineResult()
        try:
            self.open()
            conn_ok, n_conn = (True, 0)
            if self.cfg.refresh_connections:
                conn_ok, n_conn = self._refresh_connections()
            res.connections_found = n_conn

            if self.cfg.refresh_pivots:
                res.pivots_ok, res.pivots_err = self._refresh_pivots(res)

            res.success = conn_ok and res.pivots_err == 0
            if not res.success and not res.error_msg:
                parts = []
                if not conn_ok:
                    parts.append("Error actualizando conexiones")
                if res.pivots_err:
                    parts.append(f"{res.pivots_err} tabla(s) dinámica(s) fallida(s)")
                res.error_msg = "; ".join(parts)
        except Exception:
            res.error_msg = traceback.format_exc()
            self.log.error(f"Excepción:\n{res.error_msg}")
        finally:
            self.close(save=res.success and self.cfg.save_on_success)
            res.duration_s = round(time.time() - t0, 2)
        return res


# ══════════════════════════════════════════════════════════════════════════════
# FALLBACK openpyxl
# ══════════════════════════════════════════════════════════════════════════════

class OpenpyxlUpdater:
    def __init__(self, cfg: EngineConfig, log: logging.Logger):
        self.cfg = cfg
        self.log = log

    def run(self) -> EngineResult:
        t0 = time.time()
        res = EngineResult()
        self.log.warning(
            "Modo openpyxl (sin COM): conexiones externas y tablas dinámicas "
            "NO se refrescan automáticamente."
        )
        try:
            wb = openpyxl.load_workbook(self.cfg.file_path, keep_vba=True)
            self.log.info(f"Hojas: {wb.sheetnames}")
            # ── PUNTO DE EXTENSIÓN: modificar datos aquí ──
            if self.cfg.save_on_success:
                wb.save(self.cfg.file_path)
                self.log.info("Guardado (openpyxl).")
            res.success = True
        except Exception:
            res.error_msg = traceback.format_exc()
            self.log.error(res.error_msg)
        finally:
            res.duration_s = round(time.time() - t0, 2)
        return res


# ══════════════════════════════════════════════════════════════════════════════
# PUNTO DE ENTRADA PÚBLICO
# ══════════════════════════════════════════════════════════════════════════════

def run_update(cfg: EngineConfig, log: logging.Logger) -> EngineResult:
    """Valida el archivo y ejecuta el motor adecuado."""
    path = resolve_path(cfg.file_path)
    cfg.file_path = path
    log.info(f"Ruta resuelta: {path}")

    if not Path(path).exists():
        msg = f"Archivo no encontrado: {path}"
        log.error(msg)
        return EngineResult(error_msg=msg)

    # Si es un placeholder de OneDrive, disparar la descarga antes del loop
    if _is_onedrive_placeholder(path):
        log.warning(
            "El archivo está solo en la nube (OneDrive Files On-Demand). "
            "Iniciando descarga local..."
        )
        _trigger_onedrive_download(path, log)

    available = wait_for_file(path, cfg.lock_timeout, cfg.lock_retry, log, cfg.lock_max_retries, cfg.stop_event)
    if not available:
        limit = f"{cfg.lock_max_retries} intentos" if cfg.lock_max_retries else f"{cfg.lock_timeout}s"
        msg = f"Archivo bloqueado: se canceló la tarea tras {limit}."
        log.error(msg)
        return EngineResult(error_msg=msg)

    if WIN32_AVAILABLE:
        log.info("Motor: Excel COM (win32com)")
        return ExcelCOMUpdater(cfg, log).run()
    elif OPENPYXL_AVAILABLE:
        log.info("Motor: openpyxl (fallback)")
        return OpenpyxlUpdater(cfg, log).run()
    else:
        msg = "Sin motor disponible. Instala: pip install openpyxl pywin32"
        log.critical(msg)
        return EngineResult(error_msg=msg)
